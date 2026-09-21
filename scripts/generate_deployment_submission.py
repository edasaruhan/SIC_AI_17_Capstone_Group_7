"""Generate the final Turkish deployment submission and deterministic diagrams."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "academic/submissions/07_Deployment_Submission.md"
TEMPLATE = ROOT / "references/instructor/Deployment Submission.docx"
OUTPUT = ROOT / "academic/submissions/07_Deployment_Submission.docx"
DELIVERY = ROOT / "ödevler/07_Deployment/Odev_07_Model_Dagitimi_Deployment.docx"
FIGURES = ROOT / "academic/figures"

NAVY = "173B52"
BLUE = "2E6F95"
TEAL = "2D766F"
MUTED = "5B6870"
PALE = "EEF4F6"
PALE_ALT = "F7F9FA"
GRID = "D9D9D9"
WHITE = "FFFFFF"


def set_font(
    run: Any,
    *,
    name: str = "Times New Roman",
    size: float | None = None,
    color: str = "000000",
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.98)
    section.bottom_margin = Inches(0.98)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.98)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    style_names = [style.name for style in doc.styles]
    title = doc.styles["Title"] if "Title" in style_names else doc.styles.add_style("Title", 1)
    title.font.name = "Times New Roman"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 14, 15, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    style_names = [style.name for style in doc.styles]
    code = (
        doc.styles["Deployment Code"]
        if "Deployment Code" in style_names
        else doc.styles.add_style("Deployment Code", 1)
    )
    code.base_style = normal
    code.font.name = "Menlo"
    code.font.size = Pt(7.4)
    code.font.color.rgb = RGBColor.from_string(NAVY)
    code._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Menlo")
    code._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Menlo")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("GROWTHPILOT AI  |  MODEL DAĞITIMI"), size=8, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("Şahin Başcı  •  Sayfa "), size=8, color=MUTED)
    add_field(footer, "PAGE")
    set_font(footer.add_run(" / "), size=8, color=MUTED)
    add_field(footer, "NUMPAGES")


def add_inline(paragraph: Any, text: str, *, size: float | None = None) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            set_font(paragraph.add_run(part[1:-1]), name="Menlo", size=size or 9.2, color=NAVY)
        elif part.startswith("**") and part.endswith("**"):
            set_font(paragraph.add_run(part[2:-2]), size=size, bold=True)
        else:
            set_font(paragraph.add_run(part), size=size)


def set_cell(
    cell: Any,
    text: str,
    *,
    header: bool,
    alternate: bool = False,
    padding: int = 90,
) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), NAVY if header else (PALE_ALT if alternate else WHITE))
    tc_pr.append(shd)
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", padding), ("start", 110), ("bottom", padding), ("end", 110)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_inline(p, text, size=8.4 if header else 8.2)
    for run in p.runs:
        run.bold = header
        run.font.color.rgb = RGBColor.from_string(WHITE if header else "000000")


def table_widths(headers: list[str]) -> list[int]:
    if len(headers) == 2:
        return [2500, 6400]
    if headers[:2] == ["Yetenek", "Durum"]:
        return [2300, 2000, 4600]
    if len(headers) == 3:
        return [2400, 2400, 4100]
    if headers and headers[0] == "Yöntem ve yol":
        return [2000, 2850, 1600, 2450]
    return [1800, 2850, 1700, 2550]


def set_table_borders(table: Any) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), GRID)
        borders.append(border)
    table._tbl.tblPr.append(borders)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    widths = table_widths(headers)
    compact = headers[:2] == ["Yetenek", "Durum"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, text, width in zip(table.rows[0].cells, headers, widths, strict=True):
        cell.width = Pt(width / 20)
        set_cell(cell, text, header=True, padding=55 if compact else 90)
    for index, values in enumerate(rows):
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, text, width in zip(row.cells, values, widths, strict=True):
            cell.width = Pt(width / 20)
            set_cell(
                cell,
                text,
                header=False,
                alternate=index % 2 == 1,
                padding=55 if compact else 90,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_cover(doc: Document, title: str, subtitle: str, preface: list[str]) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(p.add_run(title.removesuffix(" Model Dağıtımı")), size=24, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("Model Dağıtımı"), size=18, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run(subtitle), size=13.5, color="000000")
    for line in preface:
        if line.startswith("**Yapay"):
            doc.add_paragraph().paragraph_format.space_after = Pt(5)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, line)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line.replace("  ", ""), size=11.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("DAĞITIM DURUMU"), size=9, color=TEAL, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_font(
        p.add_run(
            "Yerel uygulama ve test kanıtı mevcuttur. Canlı AWS dağıtımı, üretim OIDC, "
            "yönetilen alarm sistemi ve gerçek müşteri performans doğrulaması yoktur."
        ),
        size=11,
        color=NAVY,
        bold=True,
    )
    doc.add_page_break()


def add_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    inline = p.add_run().add_picture(str(path), width=Inches(6.15))._inline
    inline.docPr.set("title", caption.split(".", 1)[0])
    inline.docPr.set("descr", caption)


def render_markdown(doc: Document, lines: list[str], start: int) -> None:
    index = start
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
        elif stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            p = doc.add_paragraph(style="Deployment Code")
            p.paragraph_format.keep_together = True
            set_font(p.add_run("\n".join(code_lines)), name="Menlo", size=7.4, color=NAVY)
        elif stripped.startswith("| "):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("| "):
                table_lines.append(lines[index].strip())
                index += 1
            index -= 1
            parsed = [[part.strip() for part in row.strip("|").split("|")] for row in table_lines]
            add_table(doc, parsed[0], parsed[2:])
        elif stripped.startswith("!["):
            match = re.match(r"!\[(.+)]\((.+)\)", stripped)
            if not match:
                raise ValueError(stripped)
            caption, relative = match.groups()
            path = (SOURCE.parent / relative).resolve()
            add_figure(doc, path, caption)
            lookahead = index + 1
            while lookahead < len(lines) and not lines[lookahead].strip():
                lookahead += 1
            if lookahead < len(lines) and lines[lookahead].strip().startswith("*Şekil"):
                caption_p = doc.add_paragraph()
                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_p.paragraph_format.space_after = Pt(7)
                set_font(
                    caption_p.add_run(lines[lookahead].strip().strip("*")),
                    size=8.5,
                    color=MUTED,
                    italic=True,
                )
                index = lookahead
        elif stripped.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            add_inline(p, f"• {stripped[2:]}")
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.24)
            add_inline(p, stripped)
        elif stripped.startswith("*") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(stripped.strip("*")), size=8.5, color=MUTED, italic=True)
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
        index += 1


def box(ax: Any, x: float, y: float, w: float, h: float, title: str, detail: str) -> None:
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.012,rounding_size=0.018",
        facecolor="#EEF4F6",
        edgecolor="#2E6F95",
        linewidth=1.3,
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h * 0.62, title, ha="center", va="center", fontsize=10, weight="bold")
    ax.text(
        x + w / 2,
        y + h * 0.30,
        detail,
        ha="center",
        va="center",
        fontsize=7.5,
        color="#47545B",
    )


def arrow(ax: Any, start: tuple[float, float], end: tuple[float, float]) -> None:
    ax.add_patch(
        FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=12, color="#2D766F", lw=1.5)
    )


def serving_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 6.2))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    nodes = [
        (0.03, 0.61, 0.15, 0.18, "İstemci", "Arayüz veya API çağrısı"),
        (0.22, 0.61, 0.15, 0.18, "FastAPI", "İstek kimliği ve şema"),
        (0.41, 0.61, 0.16, 0.18, "Kimlik ve Kapsam", "JWT RBAC RLS"),
        (0.61, 0.61, 0.16, 0.18, "Özellik Üretimi", "Organizasyon verisi ve kesim"),
        (0.81, 0.61, 0.16, 0.18, "Dondurulmuş Model", "Manifest ve SHA-256"),
        (0.60, 0.18, 0.17, 0.18, "Karar Katmanı", "Eşik rıza insan onayı"),
        (0.80, 0.18, 0.17, 0.18, "PostgreSQL", "Tahmin denetim outbox"),
    ]
    for node in nodes:
        box(ax, *node)
    for a, b in ((0.18, 0.22), (0.37, 0.41), (0.57, 0.61), (0.77, 0.81)):
        arrow(ax, (a, 0.70), (b, 0.70))
    arrow(ax, (0.89, 0.60), (0.69, 0.37))
    arrow(ax, (0.77, 0.27), (0.80, 0.27))
    ax.text(0.5, 0.93, "GrowthPilot model sunum akışı", ha="center", fontsize=16, weight="bold")
    ax.text(
        0.5,
        0.06,
        "Model skoru pazarlama eylemi değildir  •  action_authorized = false",
        ha="center",
        fontsize=10,
        color="#173B52",
        weight="bold",
    )
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def monitoring_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 6.2))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    box(ax, 0.04, 0.64, 0.20, 0.18, "API ve Çalışan", "İstekler ve toplu işler")
    box(ax, 0.31, 0.64, 0.20, 0.18, "Uygulanan Sinyaller", "Sağlık günlük sayaç iz denetim")
    box(ax, 0.58, 0.64, 0.18, 0.18, "Toplama Sınırı", "Yetkili ölçüt uç noktası")
    box(ax, 0.80, 0.64, 0.16, 0.18, "Yerel Kanıt", "Testler ve kayıtlar")
    for a, b in ((0.24, 0.31), (0.51, 0.58), (0.76, 0.80)):
        arrow(ax, (a, 0.73), (b, 0.73))
    box(ax, 0.18, 0.20, 0.25, 0.18, "Üretim Sistem İzleme", "Hedef pano alarm nöbet")
    box(ax, 0.57, 0.20, 0.25, 0.18, "Üretim Model İzleme", "Kayma kalibrasyon gecikmeli etiket")
    arrow(ax, (0.67, 0.63), (0.33, 0.39))
    arrow(ax, (0.67, 0.63), (0.69, 0.39))
    ax.text(
        0.5,
        0.93,
        "İzleme ve üretim doğrulama sınırı",
        ha="center",
        fontsize=16,
        weight="bold",
    )
    ax.text(
        0.5,
        0.07,
        "Alt katman uygulanmadı  •  eşikler ve sorumlular canlıya geçiş öncesi tanımlanmalı",
        ha="center",
        fontsize=9.5,
        color="#8A4A24",
        weight="bold",
    )
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    serving_figure(FIGURES / "deployment_serving_architecture.png")
    monitoring_figure(FIGURES / "deployment_monitoring_flow.png")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ")
    subtitle = lines[2]
    first_section = next(i for i, line in enumerate(lines) if line.startswith("## "))
    preface = [line for line in lines[3:first_section] if line.strip()]
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Instructor template is required: {TEMPLATE}")
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    configure(doc)
    add_cover(doc, title, subtitle, preface)
    render_markdown(doc, lines, first_section)
    doc.core_properties.title = title
    doc.core_properties.subject = "Samsung Innovation Campus Model Deployment Submission"
    doc.core_properties.author = "Şahin Başcı"
    doc.core_properties.keywords = "GrowthPilot AI, model deployment, FastAPI, MLOps"
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    DELIVERY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, DELIVERY)


if __name__ == "__main__":
    main()
